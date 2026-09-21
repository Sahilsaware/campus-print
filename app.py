import os
import json
import uuid
import time
import base64
from flask import Flask, request, jsonify, render_template, session, redirect, url_for
from flask_cors import CORS
from pypdf import PdfReader
from flask_sock import Sock

app = Flask(__name__)
CORS(app)
sock = Sock(app)
app.secret_key = 'campus_print_secure_admin_key_2026'

HISTORY_FILE = 'print_history.json'
PENDING_JOBS_FILE = 'pending_jobs.json'
ADMINS_FILE = 'admins.json'
UPLOAD_FOLDER = 'uploads'

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

SUPER_ADMIN_USER = "campus_admin"
SUPER_ADMIN_PASS = "CampusPrint@2026#Secure"

last_heartbeat_time = 0
PI_PRINTER_ONLINE = False

# Global variable to track live printer state (paper status, pause state, etc.)
printer_status_global = {
    "status": "ready",
    "message": "Printer is ready"
}

def load_history():
    if os.path.exists(HISTORY_FILE):
        try:
            with open(HISTORY_FILE, 'r') as f:
                return json.load(f)
        except Exception:
            return []
    return []

def save_history(history):
    with open(HISTORY_FILE, 'w') as f:
        json.dump(history, f, indent=4)

def load_pending_jobs():
    if os.path.exists(PENDING_JOBS_FILE):
        try:
            with open(PENDING_JOBS_FILE, 'r') as f:
                return json.load(f)
        except Exception:
            return []
    return []

def save_pending_jobs(jobs):
    with open(PENDING_JOBS_FILE, 'w') as f:
        json.dump(jobs, f, indent=4)

def load_admins():
    if os.path.exists(ADMINS_FILE):
        try:
            with open(ADMINS_FILE, 'r') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_admins(admins):
    with open(ADMINS_FILE, 'w') as f:
        json.dump(admins, f, indent=4)

def get_page_dimensions(paper_size, orientation):
    # Dimensions in points (1 inch = 72 points)
    sizes = {
        'A4': (595.27, 841.89),
        'A3': (841.89, 1190.55)
    }
    width, height = sizes.get(paper_size.upper(), sizes['A4'])
    if orientation.lower() == 'landscape':
        return (height, width)
    return (width, height)

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        if username == SUPER_ADMIN_USER and password == SUPER_ADMIN_PASS:
            session['admin_logged_in'] = True
            session['username'] = username
            session['role'] = 'Super Admin'
            return redirect(url_for('admin_panel'))
        admins = load_admins()
        if username in admins and admins[username]['password'] == password:
            session['admin_logged_in'] = True
            session['username'] = username
            session['role'] = admins[username].get('role', 'Sub-Admin')
            return redirect(url_for('admin_panel'))
        return render_template('admin_login.html', error='Invalid username or password')
    return render_template('admin_login.html')

@app.route('/admin/logout')
def admin_logout():
    session.clear()
    return redirect(url_for('admin_login'))

@app.route('/admin')
def admin_panel():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
    history = load_history()
    pending_jobs = load_pending_jobs()
    return render_template('admin.html', pending_jobs=pending_jobs, history=history, username=session.get('username'), printer_status=printer_status_global)

@app.route('/count-multiple-pages', methods=['POST'])
def count_multiple_pages():
    if 'files' not in request.files:
        return jsonify({'error': 'No files uploaded'}), 400
    
    total_pages = 0
    files = request.files.getlist('files')
    for file in files:
        if file.filename != '':
            ext = os.path.splitext(file.filename)[1].lower()
            if ext == '.pdf':
                try:
                    from io import BytesIO
                    file_bytes = file.read()
                    reader = PdfReader(BytesIO(file_bytes))
                    total_pages += len(reader.pages)
                except Exception:
                    total_pages += 1
            else:
                total_pages += 1
                
    return jsonify({'success': True, 'total_pages': total_pages})

@app.route('/preview-convert', methods=['POST'])
def preview_convert():
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'Empty filename'}), 400
    
    ext = os.path.splitext(file.filename)[1].lower()
    job_id = str(uuid.uuid4())
    local_path = os.path.join(UPLOAD_FOLDER, f"{job_id}{ext}")
    file.save(local_path)
    
    print_path = local_path
    target_pagesize = get_page_dimensions('A4', 'portrait')
    
    if ext in ['.docx', '.doc']:
        try:
            import docx
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            
            doc = docx.Document(local_path)
            converted_pdf_path = os.path.join(UPLOAD_FOLDER, f"{job_id}.pdf")
            pdf_doc = SimpleDocTemplate(converted_pdf_path, pagesize=target_pagesize)
            styles = getSampleStyleSheet()
            story = []
            for element in doc.element.body:
                if element.tag.endswith('p'):
                    para = docx.text.paragraph.Paragraph(element, doc)
                    if para.text.strip():
                        story.append(Paragraph(para.text, styles['Normal']))
                        story.append(Spacer(1, 8))
            pdf_doc.build(story)
            if os.path.exists(converted_pdf_path):
                print_path = converted_pdf_path
        except Exception as e:
            return jsonify({'error': str(e)}), 500

    elif ext == '.pptx':
        try:
            from pptx import Presentation
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            
            prs = Presentation(local_path)
            converted_pdf_path = os.path.join(UPLOAD_FOLDER, f"{job_id}.pdf")
            pdf_doc = SimpleDocTemplate(converted_pdf_path, pagesize=target_pagesize)
            styles = getSampleStyleSheet()
            story = []
            for slide_idx, slide in enumerate(prs.slides):
                story.append(Paragraph(f"<b>--- Slide {slide_idx + 1} ---</b>", styles['Heading2']))
                story.append(Spacer(1, 6))
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                story.append(Paragraph(paragraph.text, styles['Normal']))
                                story.append(Spacer(1, 6))
                story.append(Spacer(1, 12))
            pdf_doc.build(story)
            if os.path.exists(converted_pdf_path):
                print_path = converted_pdf_path
        except Exception as e:
            return jsonify({'error': str(e)}), 500

    if print_path.endswith('.pdf') and os.path.exists(print_path):
        with open(print_path, 'rb') as pf:
            b64_data = base64.b64encode(pf.read()).decode('utf-8')
        return jsonify({'pdf_base64': b64_data})
    
    return jsonify({'error': 'Conversion failed'}), 500

@app.route('/print-multiple', methods=['POST'])
def print_multiple():
    if 'files' not in request.files:
        return jsonify({'error': 'No files uploaded'}), 400

    copies = int(request.form.get('copies', 1))
    orientation = request.form.get('orientation', 'portrait')
    color_mode = request.form.get('color_mode', 'bw')
    paper_size = request.form.get('paper_size', 'A4')
    duplex = request.form.get('duplex', 'false').lower() == 'true'
    page_range = request.form.get('page_range', '')
    pages_per_sheet = int(request.form.get('pages_per_sheet', 1))
    
    price_per_page = 5 if color_mode == 'color' else 2
    files = request.files.getlist('files')

    pending_jobs = load_pending_jobs()
    target_pagesize = get_page_dimensions(paper_size, orientation)

    for file in files:
        if file.filename != '':
            ext = os.path.splitext(file.filename)[1].lower()
            if ext in ['.pdf', '.png', '.jpg', '.jpeg', '.docx', '.pptx', '.doc']:
                job_id = str(uuid.uuid4())
                file_bytes = file.read()
                
                local_filename = f"{job_id}{ext}"
                local_path = os.path.join(UPLOAD_FOLDER, local_filename)
                with open(local_path, 'wb') as f:
                    f.write(file_bytes)

                print_path = local_path
                
                # Word to PDF conversion
                if ext in ['.docx', '.doc']:
                    try:
                        import docx
                        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
                        from reportlab.lib.styles import getSampleStyleSheet
                        from reportlab.lib import colors
                        
                        doc = docx.Document(local_path)
                        pdf_filename = f"{job_id}.pdf"
                        converted_pdf_path = os.path.join(UPLOAD_FOLDER, pdf_filename)
                        
                        pdf_doc = SimpleDocTemplate(converted_pdf_path, pagesize=target_pagesize)
                        styles = getSampleStyleSheet()
                        story = []
                        
                        for element in doc.element.body:
                            if element.tag.endswith('p'):
                                para = docx.text.paragraph.Paragraph(element, doc)
                                if para.text.strip():
                                    story.append(Paragraph(para.text, styles['Normal']))
                                    story.append(Spacer(1, 8))
                            elif element.tag.endswith('tbl'):
                                table = docx.table.Table(element, doc)
                                table_data = []
                                for row in table.rows:
                                    row_data = [Paragraph(cell.text.strip(), styles['Normal']) for cell in row.cells]
                                    table_data.append(row_data)
                                if table_data:
                                    t = Table(table_data)
                                    t.setStyle(TableStyle([
                                        ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
                                        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                                        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                                        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.grey),
                                        ('BOX', (0,0), (-1,-1), 1, colors.black),
                                    ]))
                                    story.append(t)
                                    story.append(Spacer(1, 10))
                                    
                        pdf_doc.build(story)
                        if os.path.exists(converted_pdf_path):
                            print_path = converted_pdf_path
                    except Exception as e:
                        print(f"Word conversion error: {e}")

                # PPT to PDF conversion
                elif ext == '.pptx':
                    try:
                        from pptx import Presentation
                        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                        from reportlab.lib.styles import getSampleStyleSheet
                        
                        prs = Presentation(local_path)
                        pdf_filename = f"{job_id}.pdf"
                        converted_pdf_path = os.path.join(UPLOAD_FOLDER, pdf_filename)
                        
                        pdf_doc = SimpleDocTemplate(converted_pdf_path, pagesize=target_pagesize)
                        styles = getSampleStyleSheet()
                        story = []
                        
                        for slide_idx, slide in enumerate(prs.slides):
                            story.append(Paragraph(f"<b>--- Slide {slide_idx + 1} ---</b>", styles['Heading2']))
                            story.append(Spacer(1, 6))
                            for shape in slide.shapes:
                                if shape.has_text_frame:
                                    for paragraph in shape.text_frame.paragraphs:
                                        if paragraph.text.strip():
                                            story.append(Paragraph(paragraph.text, styles['Normal']))
                                            story.append(Spacer(1, 6))
                            story.append(Spacer(1, 12))
                            
                        pdf_doc.build(story)
                        if os.path.exists(converted_pdf_path):
                            print_path = converted_pdf_path
                    except Exception as e:
                        print(f"PPT conversion error: {e}")

                pages = 1
                if print_path.endswith('.pdf'):
                    try:
                        reader = PdfReader(print_path)
                        pages = len(reader.pages)
                    except Exception:
                        pass
                
                total_price = pages * copies * price_per_page

                file_b64 = ""
                if os.path.exists(print_path):
                    with open(print_path, "rb") as pf:
                        file_b64 = base64.b64encode(pf.read()).decode('utf-8')

                job_data = {
                    'id': job_id,
                    'filename': file.filename,
                    'file_data': file_b64,
                    'copies': copies,
                    'orientation': orientation,
                    'color_mode': color_mode,
                    'paper_size': paper_size,
                    'duplex': duplex,
                    'page_range': page_range,
                    'pages_per_sheet': pages_per_sheet,
                    'total_price': total_price,
                    'timestamp': time.strftime('%Y-%m-%d %H:%M:%S'),
                    'status': 'pending'
                }

                pending_jobs.append(job_data)

    save_pending_jobs(pending_jobs)
    return jsonify({'success': True, 'message': 'Print job queued successfully!'})

@app.route('/get-pending-jobs', methods=['GET', 'POST'])
def get_pending_jobs():
    global last_heartbeat_time, PI_PRINTER_ONLINE
    last_heartbeat_time = time.time()
    
    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        if 'printer_online' in data:
            PI_PRINTER_ONLINE = bool(data['printer_online'])
    
    pending_jobs = load_pending_jobs()
    unlocked_jobs = [j for j in pending_jobs if j.get('status', 'pending') == 'pending']
    
    for j in unlocked_jobs:
        j['status'] = 'processing'
    save_pending_jobs(pending_jobs)
    
    return jsonify({'jobs': unlocked_jobs})

@app.route('/update-status', methods=['POST'])
def update_status():
    global PI_PRINTER_ONLINE, last_heartbeat_time
    last_heartbeat_time = time.time()
    data = request.get_json(silent=True) or {}
    PI_PRINTER_ONLINE = bool(data.get('printer_online', False))
    return jsonify({'success': True})

@app.route('/printer-status', methods=['GET'])
def printer_status():
    global last_heartbeat_time, PI_PRINTER_ONLINE, printer_status_global
    
    if time.time() - last_heartbeat_time > 15:
        PI_PRINTER_ONLINE = False
        
    response_data = {
        "online": PI_PRINTER_ONLINE,
        "status": printer_status_global.get("status", "ready"),
        "message": printer_status_global.get("message", "Printer is ready")
    }
    return jsonify(response_data)

@app.route('/update-printer-status', methods=['POST'])
def update_printer_status():
    global PI_PRINTER_ONLINE, printer_status_global, last_heartbeat_time
    last_heartbeat_time = time.time()
    data = request.get_json(silent=True) or {}
    
    if data:
        if 'printer_online' in data:
            PI_PRINTER_ONLINE = bool(data.get('printer_online'))
        
        status_text = data.get("status_message") or data.get("message", "Printer is ready")
        printer_status_global["status"] = "ready" if PI_PRINTER_ONLINE else "offline"
        printer_status_global["message"] = status_text
        
        return jsonify({"success": True})
    return jsonify({"success": False, "error": "Invalid data"}), 400

@app.route('/complete-job/<job_id>', methods=['POST'])
def complete_job(job_id):
    pending_jobs = load_pending_jobs()
    job = next((j for j in pending_jobs if j['id'] == job_id), None)
    if job:
        pending_jobs = [j for j in pending_jobs if j['id'] != job_id]
        save_pending_jobs(pending_jobs)
        history = load_history()
        history.insert(0, job)
        save_history(history)
        return jsonify({'success': True})
    return jsonify({'error': 'Job not found'}), 404

@app.route('/fail-job/<job_id>', methods=['POST'])
def fail_job(job_id):
    pending_jobs = load_pending_jobs()
    for j in pending_jobs:
        if j['id'] == job_id:
            j['status'] = 'failed'
    save_pending_jobs(pending_jobs)
    return jsonify({'success': True})

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
